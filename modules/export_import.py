from __future__ import annotations

from datetime import date
from pathlib import Path
from typing import Any

import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from . import repository as repo
from . import services
from .config import EXPORT_DIR


EXPORT_SHEETS = ["01_单词库", "02_同义词", "03_同根词", "04_复习"]

SIMPLE_COLUMNS = {
    "word": "英文单词",
    "core_meaning_cn": "中文释义",
    "part_of_speech": "词性",
    "core_meaning_en": "英文释义",
    "word_root": "词根",
    "root_key": "词根索引",
    "root_variants": "词根写法",
    "synonyms": "同义词",
    "added_date": "添加日期",
    "familiarity": "熟悉度",
    "wrong_count": "错误次数",
    "next_review_date": "下次复习日期",
}


def _dataframe_from_rows(rows: list[dict]) -> pd.DataFrame:
    return pd.DataFrame(rows) if rows else pd.DataFrame()


def _rename_simple_columns(rows: list[dict], columns: list[str]) -> pd.DataFrame:
    frame = _dataframe_from_rows(rows)
    if frame.empty:
        return frame
    return frame[columns].rename(columns=SIMPLE_COLUMNS)


def word_bank_dataframe() -> pd.DataFrame:
    return _rename_simple_columns(
        repo.simple_word_rows({"sort_by": "word", "include_synonyms": True}),
        ["word", "added_date", "core_meaning_cn", "part_of_speech", "core_meaning_en", "word_root", "synonyms"],
    )


def synonyms_dataframe() -> pd.DataFrame:
    rows = []
    for item in repo.synonym_index_rows():
        members = repo.words_for_synonym_key(item["synonym_key"])
        rows.append(
            {
                "同义索引": item["synonym_key"],
                "释义写法": item["meaning_variants"],
                "英文单词": ", ".join(member["word"] for member in members),
                "对应英文释义": "; ".join(
                    f"{member['word']}: {member['core_meaning_en']}" for member in members if member["core_meaning_en"]
                ),
                "数量": item["word_count"],
            }
        )
    return _dataframe_from_rows(rows)


def roots_dataframe() -> pd.DataFrame:
    rows = []
    for item in repo.root_index_rows():
        members = repo.words_for_root(item["root_key"])
        rows.append(
            {
                "词根索引": item["root_key"],
                "词根写法": item["root_variants"],
                "英文单词": ", ".join(member["word"] for member in members),
                "中文释义": "; ".join(f"{member['word']}: {member['core_meaning_cn']}" for member in members),
                "数量": item["word_count"],
            }
        )
    return _dataframe_from_rows(rows)


def review_dataframe() -> pd.DataFrame:
    return _rename_simple_columns(
        repo.simple_word_rows({"due_only": True, "sort_by": "review"}),
        [
            "word",
            "core_meaning_cn",
            "part_of_speech",
            "core_meaning_en",
            "word_root",
            "familiarity",
            "wrong_count",
            "next_review_date",
        ],
    )


def _append_dataframe(ws, frame: pd.DataFrame) -> None:
    if frame.empty:
        ws.append(["暂无数据"])
        return
    ws.append(list(frame.columns))
    for row in frame.itertuples(index=False, name=None):
        ws.append(list(row))


def _style_sheet(ws) -> None:
    ws.sheet_view.showGridLines = False
    if ws.max_row == 1 and ws.max_column == 1 and ws["A1"].value == "暂无数据":
        ws["A1"].font = Font(name="Microsoft YaHei", bold=True, color="555555")
        ws.column_dimensions["A"].width = 18
        return

    header_fill = PatternFill("solid", fgColor="244062")
    header_font = Font(name="Microsoft YaHei", bold=True, color="FFFFFF", size=10)
    body_font = Font(name="Microsoft YaHei", size=10)
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.font = body_font
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    for column_cells in ws.columns:
        column_letter = get_column_letter(column_cells[0].column)
        max_len = 0
        for cell in column_cells:
            value = "" if cell.value is None else str(cell.value)
            max_len = max(max_len, min(len(value), 60))
        ws.column_dimensions[column_letter].width = max(12, min(max_len + 2, 42))


def export_to_excel(output_path: Path | None = None) -> Path:
    EXPORT_DIR.mkdir(parents=True, exist_ok=True)
    if output_path is None:
        output_path = EXPORT_DIR / f"GRE_Vocabulary_{date.today().strftime('%Y%m%d')}.xlsx"

    sheets = {
        "01_单词库": word_bank_dataframe(),
        "02_同义词": synonyms_dataframe(),
        "03_同根词": roots_dataframe(),
        "04_复习": review_dataframe(),
    }

    workbook = Workbook()
    first = True
    for sheet_name, frame in sheets.items():
        ws = workbook.active if first else workbook.create_sheet(sheet_name)
        ws.title = sheet_name
        first = False
        _append_dataframe(ws, frame)
        _style_sheet(ws)

    workbook.save(output_path)
    return output_path


def verify_export(path: Path) -> dict[str, Any]:
    workbook = load_workbook(path, read_only=False, data_only=False)
    missing = [sheet for sheet in EXPORT_SHEETS if sheet not in workbook.sheetnames]
    sheet_rows = {sheet: workbook[sheet].max_row for sheet in workbook.sheetnames}
    return {"path": str(path), "missing_sheets": missing, "sheet_rows": sheet_rows}


def import_dataframe(frame: pd.DataFrame, file_name: str = "uploaded file") -> dict[str, Any]:
    rows = ["\t".join(str(value) for value in row) for row in frame.fillna("").itertuples(index=False, name=None)]
    result = services.batch_import_simple_words("\n".join(rows))
    from .db import get_connection

    with get_connection() as conn:
        conn.execute(
            """
            INSERT INTO import_logs(file_name, rows_seen, rows_created, rows_skipped, notes)
            VALUES (?, ?, ?, ?, ?)
            """,
            (
                file_name,
                result["seen"],
                len(result["created"]) + len(result["updated"]),
                len(result["skipped"]),
                "\n".join(str(item) for item in result["skipped"][:50]),
            ),
        )
        conn.commit()
    return {
        "rows_seen": result["seen"],
        "created": len(result["created"]),
        "updated": len(result["updated"]),
        "skipped": len(result["skipped"]),
        "notes": result["skipped"][:20],
    }


def import_file(file_obj: Any, file_name: str) -> dict[str, Any]:
    suffix = Path(file_name).suffix.lower()
    if suffix in {".xlsx", ".xls"}:
        frame = pd.read_excel(file_obj)
    elif suffix == ".csv":
        frame = pd.read_csv(file_obj)
    else:
        raise ValueError("仅支持 .xlsx、.xls 和 .csv 文件。")
    return import_dataframe(frame, file_name=file_name)
